"""Spike E probe 2: quantify GloBI noise for pilot taxa, probe HOSTS/FungalRoot/trait datasets. Stdlib (+openpyxl for AVONET)."""
import json, sys, urllib.request, urllib.parse, collections, csv, io, re, zipfile, time
UA = {"User-Agent": "climate-twins-spike-e/0.1 (research probe)"}
def get(url, timeout=90, rng=None):
    h = dict(UA)
    if rng: h["Range"] = rng
    with urllib.request.urlopen(urllib.request.Request(url, headers=h), timeout=timeout) as r:
        return r.status, dict(r.headers), r.read()
def J(url, **kw):
    try: return json.loads(get(url, **kw)[2])
    except Exception as e: print("  ERR", url[:100], e); return None

F = ["source_taxon_name","source_taxon_path","source_taxon_path_ranks","source_taxon_external_id","interaction_type",
     "target_taxon_name","target_taxon_path","target_taxon_path_ranks","target_taxon_external_id","study_title","study_source_citation"]
def globi(params, cap=40000, page=2000):
    rows, skip = [], 0
    while skip < cap:
        qs = urllib.parse.urlencode(dict(params, type="csv", limit=page, skip=skip)) + "".join("&field="+f for f in F)
        try: b = get("https://api.globalbioticinteractions.org/interaction?"+qs, timeout=180)[2]
        except Exception as e: print("  globi err", e); break
        recs = list(csv.DictReader(io.StringIO(b.decode("utf8","replace"))))
        rows += recs
        if len(recs) < page: break
        skip += page
    return rows
def rank(row, side, rk):
    p = row[side+"_taxon_path"].split(" | "); r = row[side+"_taxon_path_ranks"].split(" | ")
    for a, b in zip(p, r):
        if b == rk: return a
    return ""
def top(c, n=10): return c.most_common(n)

print("=== A. Monarch (Danaus plexippus): what the data says, by partner family")
rows = globi({"sourceTaxon": "Danaus plexippus"})
print("total rows (both directions as seen from sourceTaxon):", len(rows), "types:", top(collections.Counter(r["interaction_type"] for r in rows), 12))
for t in ("eats", "hasHost", "visitsFlowersOf", "interactsWith", "pollinates"):
    sub = [r for r in rows if r["interaction_type"] == t and r["source_taxon_name"].startswith("Danaus")]
    fam = collections.Counter(rank(r, "target", "family") or "(no family)" for r in sub)
    n = len(sub); ap = fam.get("Apocynaceae", 0)
    print(f"  {t}: n={n}, Apocynaceae={ap} ({100*ap/max(n,1):.0f}%), distinct target species={len({r['target_taxon_name'] for r in sub})}, top families={top(fam,6)}")
sub = [r for r in rows if r["interaction_type"] in ("eats","hasHost") and r["source_taxon_name"].startswith("Danaus")]
gen = collections.Counter(rank(r,"target","genus") for r in sub)
print("  eats/hasHost top genera:", top(gen, 12))
print("  eats/hasHost distinct Asclepias spp:", len({r['target_taxon_name'] for r in sub if rank(r,'target','genus')=='Asclepias'}))
print("  named study fields present:", sum(1 for r in rows if r["study_source_citation"]), "of", len(rows))
print("  study sources (top):", top(collections.Counter(r["study_source_citation"][:80] for r in rows), 5))

print("=== B. Oaks: consumers of Quercus (genus query) by class/order")
rows = globi({"targetTaxon": "Quercus", "interactionType": "eats"})
print("rows (consumer eats Quercus*):", len(rows), "distinct consumer names:", len({r['source_taxon_name'] for r in rows}), "distinct oak spp:", len({r['target_taxon_name'] for r in rows}))
print("  by class:", top(collections.Counter(rank(r,"source","class") or "(none)" for r in rows), 8))
print("  by order:", top(collections.Counter(rank(r,"source","order") or "(none)" for r in rows), 10))
lep = collections.Counter(r["source_taxon_name"] for r in rows if rank(r,"source","order")=="Lepidoptera")
print("  Lepidoptera consumer spp:", len(lep), "records/species top:", top(lep, 8))
print("  Lepidoptera spp with exactly 1 record:", sum(1 for v in lep.values() if v==1))
print("  study sources:", top(collections.Counter(r["study_source_citation"][:80] for r in rows), 4))
# per-species host breadth for top 25 Lepidoptera consumers
print("  host breadth (distinct plant genera in GloBI 'eats' records) for top consumers:")
for sp, n in top(lep, 25):
    rr = globi({"sourceTaxon": sp, "interactionType": "eats"}, cap=4000)
    rr = [x for x in rr if x["source_taxon_name"]==sp]
    gens = collections.Counter(rank(x,"target","genus") for x in rr)
    fams = collections.Counter(rank(x,"target","family") for x in rr)
    q = gens.get("Quercus",0)
    print(f"    {sp}: Quercus-records={n}, all eats-records={len(rr)}, distinct genera={len(gens)}, families={len(fams)}, Quercus share={100*q/max(len(rr),1):.0f}%")

print("=== C. Blacklegged tick Ixodes scapularis hosts")
rows = [r for r in globi({"sourceTaxon": "Ixodes scapularis"}) if r["source_taxon_name"].startswith("Ixodes scapularis") and r["interaction_type"] in ("hasHost","parasiteOf","ectoparasiteOf","interactsWith")]
print("host-type rows:", len(rows))
cls = collections.Counter(rank(r,"target","class") or "(none)" for r in rows)
print("  host class:", top(cls, 8))
hs = collections.Counter(r["target_taxon_name"] for r in rows)
print("  distinct host names:", len(hs), "top:", top(hs, 15))
for k in ("Odocoileus virginianus","Peromyscus leucopus","Homo sapiens","Sciurus carolinensis","Tamias striatus","Turdus migratorius","Procyon lotor"):
    print("   ", k, hs.get(k,0))
print("  study sources:", top(collections.Counter(r["study_source_citation"][:80] for r in rows), 5))

print("=== D. External-id prefixes (name matching to GBIF keys)")
allrows = []
for t in ("Danaus plexippus","Acer saccharum","Quercus robur","Turdus migratorius","Erithacus rubecula","Ixodes scapularis","Vitis vinifera"):
    rr = globi({"sourceTaxon": t}, cap=8000)
    allrows += rr
    pre = collections.Counter((r["source_taxon_external_id"] or "none").split(":")[0] for r in rr if r["source_taxon_name"]==t)
    pre_t = collections.Counter((r["target_taxon_external_id"] or "none").split(":")[0] for r in rr)
    print(f"  {t}: n={len(rr)} source-id prefixes={top(pre,4)} partner-id prefixes={top(pre_t,8)}")
pn = collections.Counter(r["target_taxon_name"] for r in allrows)
print("  partner names empty/'no name':", sum(v for k,v in pn.items() if k in ("","no name")), "of", sum(pn.values()))
# GBIF match rate on a sample of distinct partner names
import random; random.seed(1)
names = [n for n in pn if n not in ("","no name")]; random.shuffle(names); names = names[:300]
mt = collections.Counter(); rk = collections.Counter()
for n in names:
    d = J("https://api.gbif.org/v1/species/match?name="+urllib.parse.quote(n)+"&strict=false", timeout=30)
    if d: mt[d.get("matchType","?")] += 1; rk[d.get("rank","?")] += 1
print("  GBIF backbone match of 300 random partner names:", dict(mt), "ranks:", top(rk,6))

print("=== E. HOSTS (NHM Data Portal CKAN)")
d = J("https://data.nhm.ac.uk/api/3/action/package_show?id=hosts")
if d and d.get("success"):
    r = d["result"]; print("  title:", r.get("title"), "| license:", r.get("license_id"), r.get("license_title"), "| modified:", r.get("metadata_modified"))
    for x in r.get("resources", []): print("  resource:", x.get("id"), x.get("name"), x.get("format"), x.get("size"), x.get("url","")[:100])
    rid = r["resources"][0]["id"] if r.get("resources") else None
else: rid = None; print("  package_show failed", str(d)[:200])
if rid:
    for q in ("Danaus plexippus", "Quercus", "Asclepias syriaca"):
        dd = J("https://data.nhm.ac.uk/api/3/action/datastore_search?resource_id="+rid+"&q="+urllib.parse.quote(q)+"&limit=3")
        if dd and dd.get("success"):
            print(f"  q='{q}': total={dd['result'].get('total')} fields={[f['id'] for f in dd['result']['fields']][:14]}")
            for rec in dd["result"]["records"][:3]: print("    ", {k:v for k,v in rec.items() if k!='_id'})

print("=== F. FungalRoot in GBIF")
d = J("https://api.gbif.org/v1/dataset/744edc21-8dd2-474e-8a0b-b8c3d56a3c2d")
if d: print("  title:", d.get("title"), "| license:", d.get("license"), "| doi:", d.get("doi"), "| modified:", d.get("modified"))
c = J("https://api.gbif.org/v1/occurrence/count?datasetKey=744edc21-8dd2-474e-8a0b-b8c3d56a3c2d"); print("  occurrence records:", c)
for sp in ("Acer saccharum","Quercus robur","Vitis vinifera","Asclepias syriaca"):
    dd = J("https://api.gbif.org/v1/occurrence/search?datasetKey=744edc21-8dd2-474e-8a0b-b8c3d56a3c2d&scientificName="+urllib.parse.quote(sp)+"&limit=3")
    if dd:
        print(f"  {sp}: n={dd.get('count')}", [ (x.get('scientificName'), x.get('occurrenceRemarks'), x.get('verbatimScientificName')) for x in dd.get('results',[])[:2]])
        if dd.get('results'): print("    keys:", [k for k in dd['results'][0].keys() if k not in ('key','gbifID')][:40])

print("=== G. GloBI snapshot / Zenodo")
d = J("https://zenodo.org/api/records?q=%22Global%20Biotic%20Interactions%22&size=5&sort=mostrecent")
if d:
    for h in d["hits"]["hits"]:
        m = h["metadata"]; print("  ", h["id"], m.get("title","")[:90], "|", m.get("license",{}).get("id"), "|", m.get("publication_date"), "|", [(f["key"], f["size"]) for f in h.get("files",[])][:4])
d = J("https://api.github.com/repos/globalbioticinteractions/globalbioticinteractions/releases/latest")
if d: print("  github latest release:", d.get("tag_name"), d.get("published_at"))

print("=== H. Plant-pollinator networks (Web of Life / Mangal)")
try:
    b = get("https://www.web-of-life.es/get_networks.php", timeout=120)[2]; recs = json.loads(b)
    nets = collections.Counter(r["network_name"].rsplit("_",1)[0] for r in recs)
    print("  Web of Life rows:", len(recs), "network types (prefix):", top(nets, 12), "n networks:", len({r['network_name'] for r in recs}))
    pl = [r for r in recs if r["network_name"].startswith("M_PL")]
    print("  plant-pollinator networks:", len({r['network_name'] for r in pl}), "rows", len(pl))
    sp = {r["species1"] for r in pl} | {r["species2"] for r in pl}
    for k in ("Danaus plexippus","Asclepias syriaca","Apis mellifera","Bombus","Quercus"):
        print("   ", k, sum(1 for s in sp if k in s))
except Exception as e: print("  WoL ERR", e)
d = J("https://mangal.io/api/v2/dataset?count=1", timeout=60)

print("=== I. Traits: licence/size via figshare + others")
for aid, nm in ((16586228,"AVONET"), (4644424,"AmphiBIO")):
    d = J("https://api.figshare.com/v2/articles/%d" % aid)
    if d: print(f"  {nm}: title={d.get('title','')[:70]} license={d.get('license',{}).get('name')} files={[(f['name'], f['size']) for f in d.get('files',[])][:6]}")
for u in ("https://esapubs.org/archive/ecol/E095/178/", "https://esapubs.org/archive/ecol/E095/045/", "https://esapubs.org/archive/ecol/E095/045/suppl-1.php",
          "https://api.github.com/repos/cropmodels/Recocrop/contents/", "https://api.github.com/repos/cropmodels/Recocrop/contents/inst"):
    try:
        s,h,b = get(u, timeout=60); t = re.sub(r"<[^>]+>", " ", b.decode("utf8","replace")); t = re.sub(r"\s+"," ",t)
        print("  ", u, s, t[:700] if "api.github" not in u else b[:900].decode())
    except Exception as e: print("  ", u, "ERR", e)
d = J("https://api.crossref.org/works?query.bibliographic=COMBINE+coalesced+mammal+database+Soria&rows=2&select=DOI,title,link")
if d: print("  crossref COMBINE:", [(i.get("DOI"), i.get("title")) for i in d["message"]["items"]])

print("=== J. AVONET rows for pilot birds")
import subprocess, os
os.makedirs("av", exist_ok=True)
subprocess.run(["pip","install","-q","openpyxl"], check=False)
try:
    b = get("https://ndownloader.figshare.com/files/34480856", timeout=300)[2]
    open("av/avonet.xlsx","wb").write(b)
    import openpyxl
    wb = openpyxl.load_workbook("av/avonet.xlsx", read_only=True)
    print("  sheets:", wb.sheetnames)
    for sn in wb.sheetnames:
        ws = wb[sn]; it = ws.iter_rows(values_only=True); hdr = next(it)
        print("  sheet", sn, "cols:", hdr[:40])
        if "AVONET1" in sn or "eBird" in sn:
            idx = {h:i for i,h in enumerate(hdr)}; n=0
            for row in it:
                n += 1
                if row[idx["Species1"]] in ("Turdus migratorius","Erithacus rubecula"):
                    print("   ", {k: row[idx[k]] for k in hdr if k in ("Species1","Mass","Hand-Wing.Index","Habitat","Migration","Trophic.Level","Primary.Lifestyle","Range.Size")})
            print("   rows:", n)
except Exception as e: print("  AVONET ERR", e)
